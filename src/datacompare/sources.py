"""数据接入层：CSV / 自定义分隔符文本 / Excel / Parquet / JSON → DuckDB。

设计要点
--------
1. **原始文本原样保留**：所有列一律以 VARCHAR 落库，
   这样 ``1.50`` 和 ``1234`` 这类「原始写法」不会在入库时就被抹掉，
   后面才能区分「值不同」和「只是格式不同」。
2. **空字符串与 NULL 分离**：DuckDB 默认把空字段读成 NULL，
   这里通过一个几乎不可能出现的哨兵值 ``nullstr`` 关掉该行为，
   空值判定完全交给 :mod:`datacompare.normalize` 的 null token 规则。
3. **大文件走磁盘**：CSV 直接由 DuckDB 流式读取，不经过 Python 内存。
   Excel 没有流式读取能力，先落临时 UTF-8 CSV 再交给 DuckDB。
"""

from __future__ import annotations

import csv
import glob as globlib
import math
import os
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .config import SourceSpec
from .sqlutil import quote_ident, sql_str

#: 用来「关闭」DuckDB 内置 null 识别的哨兵值（DuckDB 要求 nullstr 非空）
NULL_SENTINEL = "@@DC_NO_NULL@@"

_UTF8_ALIASES = {"utf-8", "utf8", "UTF-8", "UTF8", "u8"}
_UTF8_SIG_ALIASES = {"utf-8-sig", "utf8-sig", "UTF-8-SIG"}

_DUP_SUFFIX_RE = None


@dataclass
class LoadedSource:
    """一侧数据加载完成后的元信息。"""

    table: str
    columns: List[str]
    row_count: int
    path: str
    kind: str
    temp_paths: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def ncols(self) -> int:
        return len(self.columns)


# --------------------------------------------------------------------------
# 对外入口
# --------------------------------------------------------------------------
def load_source(con, spec: SourceSpec, table: str, temp_dir: Optional[str] = None) -> LoadedSource:
    """把一侧数据加载成 DuckDB 表（全 VARCHAR + 原始顺序保留）。"""
    if not spec.path:
        raise ValueError("数据源未指定 path")
    kind = spec.resolved_kind()
    if kind == "excel":
        return _load_excel(con, spec, table, temp_dir)
    if kind == "parquet":
        return _load_parquet(con, spec, table)
    if kind == "json":
        return _load_json(con, spec, table)
    return _load_text(con, spec, table, temp_dir)


ROW_ID = "__row_id"


def add_row_id(con, table: str, out_table: Optional[str] = None) -> str:
    """给表加上稳定的自增行号 ``__row_id``（从 0 开始，按源文件物理顺序）。

    ``__row_id`` 是行对齐的兜底依据，也是报告里定位问题行的坐标。
    """
    out = out_table or table
    cols = table_columns(con, table)
    if ROW_ID in cols:
        raise ValueError(
            f"数据中存在名为 {ROW_ID} 的列，与工具内部列冲突，请在配置中重命名该列"
        )
    select = f"SELECT row_number() OVER () - 1 AS {quote_ident(ROW_ID)}, * FROM "
    if out == table:
        tmp = f"__rid_{table}"
        con.execute(f"CREATE OR REPLACE TABLE {quote_ident(tmp)} AS {select}{quote_ident(table)}")
        con.execute(f"DROP TABLE {quote_ident(table)}")
        con.execute(f"ALTER TABLE {quote_ident(tmp)} RENAME TO {quote_ident(table)}")
    else:
        con.execute(f"CREATE OR REPLACE TABLE {quote_ident(out)} AS {select}{quote_ident(table)}")
    return out


def table_columns(con, table: str) -> List[str]:
    """读取表的列名（按定义顺序）。"""
    rows = con.sql(f"DESCRIBE {quote_ident(table)}").fetchall()
    return [r[0] for r in rows]


def table_row_count(con, table: str) -> int:
    return int(con.sql(f"SELECT count(*) FROM {quote_ident(table)}").fetchone()[0])


# --------------------------------------------------------------------------
# 文本 / CSV
# --------------------------------------------------------------------------
def _load_text(con, spec: SourceSpec, table: str, temp_dir: Optional[str]) -> LoadedSource:
    temp_paths: List[str] = []
    warnings: List[str] = []

    if not any(ch in (spec.path or "") for ch in "*?["):
        if not os.path.exists(spec.path):
            raise FileNotFoundError(f"找不到数据文件：{spec.path}")

    path, is_temp = _prepare_text_path(spec, temp_dir)
    if is_temp:
        temp_paths.append(path)

    options: List[str] = [
        f"header={'true' if spec.header else 'false'}",
        "all_varchar=true",
        f"nullstr=[{sql_str(NULL_SENTINEL)}]",
        f"sample_size={int(spec.sample_size)}",
    ]
    if spec.skip_rows:
        options.append(f"skip={int(spec.skip_rows)}")
    delim = spec.resolved_delimiter()
    if delim is not None:
        options.append(f"delim={sql_str(delim)}")
    if spec.quote and spec.quote != '"':
        options.append(f"quote={sql_str(spec.quote)}")
    if spec.escape:
        options.append(f"escape={sql_str(spec.escape)}")

    opts = ", ".join(options)
    try:
        con.execute(
            f"CREATE OR REPLACE TABLE {quote_ident('__raw_' + table)} AS "
            f"SELECT * FROM read_csv({sql_str(path)}, {opts})"
        )
    except Exception as exc:
        _cleanup_paths(temp_paths)
        raise RuntimeError(
            f"读取文本文件失败：{spec.path}\n"
            f"  解析参数：{opts}\n"
            f"  提示：如果是自定义分隔符/编码问题，请在配置里显式指定 "
            f"delimiter / encoding / quote。\n"
            f"  原始错误：{exc}"
        ) from exc

    columns = table_columns(con, "__raw_" + table)
    _check_duplicate_headers(columns, warnings)
    add_row_id(con, "__raw_" + table, table)
    con.execute(f"DROP TABLE IF EXISTS {quote_ident('__raw_' + table)}")

    return LoadedSource(
        table=table,
        columns=columns,
        row_count=table_row_count(con, table),
        path=spec.path,
        kind="csv",
        temp_paths=temp_paths,
        warnings=warnings,
    )


def _prepare_text_path(spec: SourceSpec, temp_dir: Optional[str]) -> Tuple[str, bool]:
    """返回可供 DuckDB 读取的 UTF-8 文件路径。

    UTF-8 文件直接用原路径；其它编码（GBK/GB18030/UTF-16 等）先转码到临时文件。
    """
    encoding = (spec.encoding or "utf-8").strip()
    if encoding in _UTF8_ALIASES:
        return spec.path, False
    if encoding in _UTF8_SIG_ALIASES:
        return _transcode_to_utf8(spec.path, encoding, temp_dir), True
    return _transcode_to_utf8(spec.path, encoding, temp_dir), True


def _transcode_to_utf8(path: str, encoding: str, temp_dir: Optional[str]) -> str:
    fd, tmp_path = tempfile.mkstemp(suffix=".utf8.csv", dir=temp_dir, prefix="dcconv_")
    os.close(fd)
    try:
        with open(path, "r", encoding=encoding, errors="replace", newline="") as src, \
                open(tmp_path, "w", encoding="utf-8", newline="") as dst:
            while True:
                chunk = src.read(1 << 20)
                if not chunk:
                    break
                dst.write(chunk)
    except LookupError as exc:
        raise ValueError(f"不认识的编码 {encoding!r}：{exc}") from exc
    return tmp_path


def _check_duplicate_headers(columns: Sequence[str], warnings: List[str]) -> None:
    """DuckDB 会把重名列自动改写成 ``name_1``，这里提示用户注意。"""
    names = set(columns)
    for col in columns:
        if "_" not in col:
            continue
        base, _, suffix = col.rpartition("_")
        if base and suffix.isdigit() and base in names:
            warnings.append(
                f"检测到重复列名：`{base}` 与 `{col}` 同时存在，"
                "DuckDB 已自动重命名，请确认列映射是否正确"
            )


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def _load_excel(con, spec: SourceSpec, table: str, temp_dir: Optional[str]) -> LoadedSource:
    if spec.path.lower().endswith(".xls"):
        raise ValueError(
            "不支持老版 .xls 格式，请先另存为 .xlsx 或导出为 CSV"
        )

    warnings: List[str] = []
    temp_paths: List[str] = []
    files = _expand_glob(spec.path)
    if not files:
        raise FileNotFoundError(f"找不到 Excel 文件：{spec.path}")

    fd, tmp_csv = tempfile.mkstemp(suffix=".csv", dir=temp_dir, prefix="dcexcel_")
    os.close(fd)
    temp_paths.append(tmp_csv)

    try:
        header: List[str] = []
        with open(tmp_csv, "w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh, lineterminator="\n")
            first = True
            for file_path in files:
                rows_iter = _iter_excel_rows(file_path, spec, warnings)
                try:
                    cols = next(rows_iter)
                except StopIteration:
                    cols = []
                if not cols:
                    warnings.append(f"{os.path.basename(file_path)} 没有读取到任何列，已跳过")
                    continue
                if first:
                    header = cols
                    writer.writerow(header)
                    first = False
                elif cols != header:
                    warnings.append(
                        f"{os.path.basename(file_path)} 的表头与首个文件不一致，已按列序对齐"
                    )
                for row in rows_iter:
                    writer.writerow(_fit_row(row, len(header)))

        if not header:
            raise ValueError(f"Excel 文件没有可用数据：{spec.path}")

        con.execute(
            f"CREATE OR REPLACE TABLE {quote_ident(table)} AS "
            f"SELECT * FROM read_csv({sql_str(tmp_csv)}, header=true, all_varchar=true, "
            f"nullstr=[{sql_str(NULL_SENTINEL)}], sample_size=-1)"
        )
    except BaseException:
        _cleanup_paths(temp_paths)
        raise

    columns = table_columns(con, table)
    if len(columns) != len(header):
        warnings.append("Excel 表头列数与实际解析列数不一致，以 DuckDB 解析结果为准")
    _check_duplicate_headers(columns, warnings)
    add_row_id(con, table, table)
    return LoadedSource(
        table=table,
        columns=columns,
        row_count=table_row_count(con, table),
        path=spec.path,
        kind="excel",
        temp_paths=temp_paths,
        warnings=warnings,
    )


def _iter_excel_rows(
    path: str, spec: SourceSpec, warnings: List[str]
) -> Iterable[List[Any]]:
    """生成器：第一个 yield 是表头，其余每次 yield 是一行数据。

    注意必须用生成器——工作簿要在消费者读完所有行之后才关闭，
    否则 read_only 模式下迭代到中途会报错。
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "读取 Excel 需要 openpyxl，请执行 pip install openpyxl"
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _pick_sheet(wb, spec.sheet, path, warnings)
        iterator = ws.iter_rows(values_only=True)
        for _ in range(int(spec.skip_rows or 0)):
            next(iterator, None)

        header_row = None
        if spec.header:
            header_row = next(iterator, None)
            while header_row is not None and _is_blank_row(header_row):
                header_row = next(iterator, None)

        # 先探一行数据，用来确定列数（表头可能比数据行短）
        first_data = next(iterator, None)
        while first_data is not None and _is_blank_row(first_data):
            first_data = next(iterator, None)

        width = 0
        if header_row:
            width = len(header_row)
        if first_data:
            width = max(width, len(first_data))
        if width == 0:
            yield []
            return

        if header_row:
            yield _make_header(header_row, width)
        else:
            yield [f"column{i}" for i in range(width)]

        if first_data is not None:
            yield [_cell_text(v) for v in first_data]
            for row in iterator:
                if row is None or _is_blank_row(row):
                    continue
                yield [_cell_text(v) for v in row]
    finally:
        wb.close()


def _cleanup_paths(paths: Sequence[str]) -> None:
    for path in paths:
        try:
            os.unlink(path)
        except OSError:
            pass


def _pick_sheet(wb, sheet: Optional[Any], path: str, warnings: List[str]):
    if sheet is None:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            raise ValueError(f"Excel 文件没有工作表：{path}")
        if len(wb.worksheets) > 1:
            warnings.append(
                f"{os.path.basename(path)} 有 {len(wb.worksheets)} 个工作表，"
                f"默认使用第一个 `{ws.title}`"
            )
        return ws
    if isinstance(sheet, int):
        try:
            return wb.worksheets[sheet]
        except IndexError as exc:
            raise ValueError(f"工作表序号 {sheet} 超出范围：{path}") from exc
    if sheet in wb.sheetnames:
        return wb[sheet]
    raise ValueError(f"找不到工作表 {sheet!r}：{path}，可用：{wb.sheetnames}")


def _make_header(row: Sequence[Any], width: int) -> List[str]:
    names: List[str] = []
    used: Dict[str, int] = {}
    for i in range(width):
        raw = row[i] if i < len(row) else None
        name = _cell_text(raw) or f"column{i}"
        name = str(name).strip() or f"column{i}"
        if name in used:
            used[name] += 1
            name = f"{name}_{used[name]}"
        else:
            used[name] = 0
        names.append(name)
    return names


def _is_blank_row(row: Sequence[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def _fit_row(row: Sequence[Any], width: int) -> List[Any]:
    values = list(row)
    if len(values) < width:
        values.extend([None] * (width - len(values)))
    elif len(values) > width:
        values = values[:width]
    return values


def _cell_text(value: Any) -> Optional[str]:
    """把 Excel 单元格转成尽量贴近原始写法的文本。"""
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if math.isinf(value):
            return "inf" if value > 0 else "-inf"
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, Decimal):
        text = format(value, "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".") or "0"
        return text
    return str(value)


# --------------------------------------------------------------------------
# Parquet / JSON
# --------------------------------------------------------------------------
def _load_parquet(con, spec: SourceSpec, table: str) -> LoadedSource:
    paths = _expand_glob(spec.path) or [spec.path]
    if len(paths) > 1:
        source = f"read_parquet({sql_str(spec.path)}, union_by_name=true)"
    else:
        source = f"read_parquet({sql_str(paths[0])})"
    _create_varchar_table(con, source, table)
    return LoadedSource(
        table=table,
        columns=table_columns(con, table),
        row_count=table_row_count(con, table),
        path=spec.path,
        kind="parquet",
    )


def _load_json(con, spec: SourceSpec, table: str) -> LoadedSource:
    paths = _expand_glob(spec.path) or [spec.path]
    if len(paths) > 1:
        source = f"read_json_auto({sql_str(spec.path)})"
    else:
        source = f"read_json_auto({sql_str(paths[0])})"
    try:
        _create_varchar_table(con, source, table)
    except Exception as exc:
        raise RuntimeError(
            f"读取 JSON 失败：{spec.path}\n"
            f"  提示：需要 DuckDB 的 json 扩展；嵌套结构会被转成文本。\n"
            f"  原始错误：{exc}"
        ) from exc
    return LoadedSource(
        table=table,
        columns=table_columns(con, table),
        row_count=table_row_count(con, table),
        path=spec.path,
        kind="json",
    )


def _create_varchar_table(con, source_sql: str, table: str) -> None:
    """把任意来源统一转成「全 VARCHAR + 带 __row_id」的表。"""
    probe = "__probe_" + table
    con.execute(f"CREATE OR REPLACE TABLE {quote_ident(probe)} AS SELECT * FROM {source_sql}")
    try:
        cols = table_columns(con, probe)
        if not cols:
            raise ValueError("数据源没有解析出任何列")
        select_list = ", ".join(
            f"TRY_CAST({quote_ident(c)} AS VARCHAR) AS {quote_ident(c)}" for c in cols
        )
        con.execute(
            f"CREATE OR REPLACE TABLE {quote_ident(table)} AS "
            f"SELECT {select_list} FROM {quote_ident(probe)}"
        )
    finally:
        con.execute(f"DROP TABLE IF EXISTS {quote_ident(probe)}")
    add_row_id(con, table, table)


def _expand_glob(pattern: str) -> List[str]:
    if not pattern:
        return []
    if any(ch in pattern for ch in "*?["):
        return sorted(globlib.glob(pattern))
    return [pattern] if os.path.exists(pattern) else []
